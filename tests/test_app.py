"""Comprehensive test suite for AI-Based Vehicle Damage Detection System.
Covers Authentication, Admin Workflow, User Workflow, Data Separation, Temporary Analysis, and Regressions.
"""
import io
import os
import shutil
from pathlib import Path
import cv2
import numpy as np
import openpyxl
import pytest

import sys
from pathlib import Path
BACKEND_DIR = Path(__file__).resolve().parent.parent / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import config
from app import app
from services import db, charts, damage_detector


@pytest.fixture(scope="function")
def test_env(tmp_path):
    """Sets up an isolated test environment with temporary database and upload dirs."""
    test_db = tmp_path / "test_app.db"
    test_uploads = tmp_path / "uploads"
    test_temp = test_uploads / "temp"
    test_data = tmp_path / "data"

    test_uploads.mkdir(parents=True, exist_ok=True)
    test_temp.mkdir(parents=True, exist_ok=True)
    test_data.mkdir(parents=True, exist_ok=True)

    # Save original paths
    orig_db = config.DB_PATH
    orig_up = config.UPLOAD_DIR
    orig_tmp = config.TEMP_UPLOAD_DIR
    orig_data = config.DATA_DIR

    config.DB_PATH = test_db
    config.UPLOAD_DIR = test_uploads
    config.TEMP_UPLOAD_DIR = test_temp
    config.DATA_DIR = test_data

    # Initialize isolated database
    db.init_db(test_db)

    # Create dummy vehicle test image
    test_img_path = tmp_path / "test_car.jpg"
    img = np.full((300, 400, 3), 180, dtype=np.uint8)
    # Draw simple car shape
    cv2.rectangle(img, (50, 100), (350, 240), (50, 50, 50), -1)
    cv2.circle(img, (100, 240), 30, (20, 20, 20), -1)
    cv2.circle(img, (300, 240), 30, (20, 20, 20), -1)
    # Add high-contrast edge simulating scratch
    cv2.line(img, (120, 150), (220, 155), (255, 255, 255), 4)
    cv2.imwrite(str(test_img_path), img)

    app.config["TESTING"] = True
    client = app.test_client()

    yield {
        "client": client,
        "test_img_path": test_img_path,
        "tmp_path": tmp_path
    }

    # Restore paths
    config.DB_PATH = orig_db
    config.UPLOAD_DIR = orig_up
    config.TEMP_UPLOAD_DIR = orig_tmp
    config.DATA_DIR = orig_data


# Helper functions
def login_admin(client):
    return client.post("/api/admin/login", json={
        "email": "admin@Samrat.com",
        "password": "Samrat@2026"
    })


def register_user(client, username="alice", email="alice@test.com", password="password123", full_name=None):
    fname = full_name or username.capitalize()
    return client.post("/api/register", json={
        "full_name": fname,
        "username": username,
        "email": email,
        "password": password,
        "confirm_password": password
    })


def login_user(client, email="alice@test.com", password="password123"):
    return client.post("/api/login", json={
        "email": email,
        "password": password
    })


# ---------------- 1. AUTHENTICATION TESTS ----------------

def test_01_admin_login_success(test_env):
    client = test_env["client"]
    res = login_admin(client)
    assert res.status_code == 200
    assert res.json.get("ok") is True
    assert res.json.get("role") == "admin"


def test_02_invalid_admin_credentials_rejected(test_env):
    client = test_env["client"]
    res = client.post("/api/admin/login", json={
        "email": "admin@Samrat.com",
        "password": "WrongPassword123"
    })
    assert res.status_code == 401
    assert "Invalid" in res.json.get("error", "")


def test_03_regular_user_cannot_access_admin_routes(test_env):
    client = test_env["client"]
    register_user(client, username="bob", email="bob@test.com")
    # Bob is logged in as regular user
    res = client.get("/api/admin/users")
    assert res.status_code == 403
    assert "Administrator privileges required" in res.json.get("error", "")

    # HTML page redirect
    res_page = client.get("/admin/dashboard")
    assert res_page.status_code == 302
    assert "/user/dashboard" in res_page.headers.get("Location", "")


def test_04_admin_registration_not_publicly_available(test_env):
    client = test_env["client"]
    # Attempt to pass role="admin" through public registration endpoint
    res = client.post("/api/register", json={
        "full_name": "Hacker Admin",
        "username": "hacker_admin",
        "email": "hacker@test.com",
        "password": "password123",
        "confirm_password": "password123",
        "role": "admin"
    })
    assert res.status_code == 200
    # Verify in DB that role was forced to 'user'
    u = db.get_user_by_email("hacker@test.com")
    assert u is not None
    assert u["role"] == "user"


def test_05_new_user_can_register_successfully(test_env):
    client = test_env["client"]
    res = register_user(client, username="clara", email="clara@test.com", full_name="Clara Oswald")
    assert res.status_code == 200
    assert res.json.get("ok") is True
    u = db.get_user_by_email("clara@test.com")
    assert u["username"] == "clara"
    assert u["full_name"] == "Clara Oswald"


def test_06_duplicate_email_registration_rejected(test_env):
    client = test_env["client"]
    register_user(client, username="dave1", email="dave@test.com")
    res = register_user(client, username="dave2", email="dave@test.com")
    assert res.status_code in (400, 409)
    assert "already registered" in res.json.get("error", "").lower()


def test_07_duplicate_username_registration_rejected(test_env):
    client = test_env["client"]
    register_user(client, username="unique_user", email="user1@test.com")
    res = register_user(client, username="unique_user", email="user2@test.com")
    assert res.status_code in (400, 409)
    assert "username is already taken" in res.json.get("error", "").lower()


def test_08_registered_user_can_login_using_email_and_password(test_env):
    client = test_env["client"]
    register_user(client, username="emma", email="emma@test.com", password="mypassword")
    client.post("/api/logout")  # Log out

    res = login_user(client, email="emma@test.com", password="mypassword")
    assert res.status_code == 200
    assert res.json.get("ok") is True
    assert res.json.get("role") == "user"


def test_09_invalid_user_login_rejected(test_env):
    client = test_env["client"]
    register_user(client, username="felix", email="felix@test.com", password="correctpass")
    client.post("/api/logout")

    res = login_user(client, email="felix@test.com", password="wrongpass")
    assert res.status_code == 401
    assert "Invalid email or password" in res.json.get("error", "")


def test_10_logout_invalidates_session(test_env):
    client = test_env["client"]
    login_admin(client)
    res_out = client.post("/api/logout")
    assert res_out.status_code == 200

    # Next call should be unauthenticated
    res_check = client.get("/api/admin/users")
    assert res_check.status_code == 401


# ---------------- 2. ADMIN WORKFLOW TESTS ----------------

def test_11_admin_can_view_dashboard(test_env):
    client = test_env["client"]
    login_admin(client)
    res = client.get("/admin/dashboard")
    assert res.status_code == 200
    assert b"Administrator Dashboard" in res.data


def test_12_admin_can_list_registered_users(test_env):
    client = test_env["client"]
    register_user(client, username="george", email="george@test.com")
    login_admin(client)

    res = client.get("/api/admin/users")
    assert res.status_code == 200
    users = res.json.get("users", [])
    assert any(u["email"] == "george@test.com" for u in users)


def test_13_admin_can_select_user_for_analysis(test_env):
    client = test_env["client"]
    register_user(client, username="helen", email="helen@test.com")
    login_admin(client)

    res = client.get("/api/admin/users/selectable")
    assert res.status_code == 200
    users = res.json.get("users", [])
    assert any(u["username"] == "helen" for u in users)


def test_14_admin_cannot_save_official_report_without_user(test_env):
    client = test_env["client"]
    login_admin(client)
    test_img = test_env["test_img_path"]

    with open(test_img, "rb") as f:
        res = client.post("/api/admin/detect-damage", data={
            "image": (f, "test_car.jpg")
            # user_id is missing!
        }, content_type="multipart/form-data")

    assert res.status_code == 400
    assert "select a registered user" in res.json.get("error", "").lower()


def test_15_admin_analysis_creates_official_report_for_selected_user(test_env):
    client = test_env["client"]
    register_user(client, username="ian", email="ian@test.com")
    user_ian = db.get_user_by_email("ian@test.com")

    login_admin(client)
    test_img = test_env["test_img_path"]

    with open(test_img, "rb") as f:
        res = client.post("/api/admin/detect-damage", data={
            "user_id": user_ian["user_id"],
            "image": (f, "test_car.jpg")
        }, content_type="multipart/form-data")

    assert res.status_code == 200
    data = res.json
    assert "id" in data
    assert data["id"].startswith("R-")

    # Verify in DB
    rep = db.get_report(data["id"])
    assert rep is not None
    assert rep["user_id"] == user_ian["user_id"]
    assert rep["created_by_admin_id"] == 1001


def test_16_saved_report_appears_in_admin_portal(test_env):
    client = test_env["client"]
    register_user(client, username="julia", email="julia@test.com")
    uid = db.get_user_by_email("julia@test.com")["user_id"]
    login_admin(client)

    with open(test_env["test_img_path"], "rb") as f:
        rep_res = client.post("/api/admin/detect-damage", data={
            "user_id": uid, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")
    rid = rep_res.json["id"]

    res = client.get("/api/admin/reports")
    assert res.status_code == 200
    reports = res.json.get("reports", [])
    assert any(r["report_id"] == rid for r in reports)


def test_17_saved_report_appears_in_selected_user_portal(test_env):
    client = test_env["client"]
    register_user(client, username="kevin", email="kevin@test.com")
    uid = db.get_user_by_email("kevin@test.com")["user_id"]

    login_admin(client)
    with open(test_env["test_img_path"], "rb") as f:
        rep_res = client.post("/api/admin/detect-damage", data={
            "user_id": uid, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")
    rid = rep_res.json["id"]
    client.post("/api/logout")

    # Kevin logs in
    login_user(client, email="kevin@test.com")
    res = client.get("/api/user/reports")
    assert res.status_code == 200
    user_reports = res.json.get("reports", [])
    assert any(r["report_id"] == rid for r in user_reports)


def test_18_report_does_not_appear_in_another_user_portal(test_env):
    client = test_env["client"]
    register_user(client, username="laura", email="laura@test.com")
    register_user(client, username="mark", email="mark@test.com")
    uid_laura = db.get_user_by_email("laura@test.com")["user_id"]

    login_admin(client)
    with open(test_env["test_img_path"], "rb") as f:
        rep_res = client.post("/api/admin/detect-damage", data={
            "user_id": uid_laura, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")
    rid = rep_res.json["id"]
    client.post("/api/logout")

    # Mark logs in
    login_user(client, email="mark@test.com")
    res = client.get("/api/user/reports")
    assert res.status_code == 200
    mark_reports = res.json.get("reports", [])
    # Mark should NOT see Laura's report!
    assert not any(r["report_id"] == rid for r in mark_reports)


def test_19_admin_can_view_login_history(test_env):
    client = test_env["client"]
    login_admin(client)
    res = client.get("/api/admin/login-history")
    assert res.status_code == 200
    history = res.json.get("history", [])
    assert len(history) > 0


def test_20_21_admin_download_excel_and_sheets(test_env):
    client = test_env["client"]
    register_user(client, username="nina", email="nina@test.com")
    uid = db.get_user_by_email("nina@test.com")["user_id"]
    login_admin(client)

    with open(test_env["test_img_path"], "rb") as f:
        client.post("/api/admin/detect-damage", data={
            "user_id": uid, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")

    res = client.get("/api/admin/download-excel")
    assert res.status_code == 200
    assert res.headers.get("Content-Type") == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # Verify workbook structure
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    assert "Reports" in wb.sheetnames
    assert "Users" in wb.sheetnames
    assert "LoginHistory" in wb.sheetnames

    # Check Reports headers
    rep_headers = [c.value for c in wb["Reports"][1]]
    assert "Report ID" in rep_headers
    assert "User ID" in rep_headers
    assert "Estimated Repair Cost Range" in rep_headers


# ---------------- 3. USER WORKFLOW TESTS ----------------

def test_22_user_can_view_own_dashboard(test_env):
    client = test_env["client"]
    register_user(client, username="oscar", email="oscar@test.com")
    res = client.get("/user/dashboard")
    assert res.status_code == 200
    assert b"Welcome back, Oscar!" in res.data


def test_23_user_can_view_own_official_reports(test_env):
    client = test_env["client"]
    register_user(client, username="paula", email="paula@test.com")
    uid = db.get_user_by_email("paula@test.com")["user_id"]

    login_admin(client)
    with open(test_env["test_img_path"], "rb") as f:
        rep_res = client.post("/api/admin/detect-damage", data={
            "user_id": uid, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")
    rid = rep_res.json["id"]
    client.post("/api/logout")

    login_user(client, email="paula@test.com")
    res = client.get(f"/api/user/reports/{rid}")
    assert res.status_code == 200
    assert res.json["report"]["report_id"] == rid


def test_24_user_cannot_view_another_user_report_by_id(test_env):
    client = test_env["client"]
    register_user(client, username="quinn", email="quinn@test.com")
    register_user(client, username="rachel", email="rachel@test.com")
    uid_quinn = db.get_user_by_email("quinn@test.com")["user_id"]

    login_admin(client)
    with open(test_env["test_img_path"], "rb") as f:
        rep_res = client.post("/api/admin/detect-damage", data={
            "user_id": uid_quinn, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")
    rid = rep_res.json["id"]
    client.post("/api/logout")

    # Rachel logs in and tries to access Quinn's report
    login_user(client, email="rachel@test.com")
    res = client.get(f"/api/user/reports/{rid}")
    assert res.status_code == 403
    assert "Access denied" in res.json.get("error", "")

    # HTML page should return 403
    res_page = client.get(f"/user/reports/{rid}")
    assert res_page.status_code == 403


def test_25_user_cannot_view_another_user_annotated_image(test_env):
    client = test_env["client"]
    register_user(client, username="sam", email="sam@test.com")
    register_user(client, username="tina", email="tina@test.com")
    uid_sam = db.get_user_by_email("sam@test.com")["user_id"]

    login_admin(client)
    with open(test_env["test_img_path"], "rb") as f:
        rep_res = client.post("/api/admin/detect-damage", data={
            "user_id": uid_sam, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")
    ann_img = rep_res.json["annotated"]
    client.post("/api/logout")

    # Tina logs in and tries to download Sam's image
    login_user(client, email="tina@test.com")
    res = client.get(f"/uploads/{ann_img}")
    assert res.status_code == 403
    assert "Access denied" in res.json.get("error", "")


def test_26_user_cannot_access_admin_analytics_or_login_history(test_env):
    client = test_env["client"]
    register_user(client, username="uma", email="uma@test.com")
    res_analytics = client.get("/api/admin/analytics/summary")
    assert res_analytics.status_code == 403

    res_logins = client.get("/api/admin/login-history")
    assert res_logins.status_code == 403


def test_27_user_cannot_download_complete_excel_workbook(test_env):
    client = test_env["client"]
    register_user(client, username="victor", email="victor@test.com")
    res = client.get("/api/admin/download-excel")
    assert res.status_code == 403

    # Compatibility endpoint also blocks regular users
    res_compat = client.get("/api/download-excel")
    assert res_compat.status_code == 403


def test_28_29_30_31_temporary_user_analysis_isolation(test_env):
    client = test_env["client"]
    register_user(client, username="wendy", email="wendy@test.com")

    # Count official reports before
    initial_reports_count = len(db.get_all_reports())

    with open(test_env["test_img_path"], "rb") as f:
        res = client.post("/api/user/analyze-temporary", data={
            "image": (f, "temp_scan.jpg")
        }, content_type="multipart/form-data")

    assert res.status_code == 200
    data = res.json
    assert data.get("is_temporary") is True
    assert "temp/" in data.get("annotated", "")

    # 30: Temporary analysis does NOT create an official report!
    new_reports_count = len(db.get_all_reports())
    assert new_reports_count == initial_reports_count

    # 31: Temporary analysis does NOT appear in user's official reports
    u_reports = client.get("/api/user/reports").json.get("reports", [])
    assert len(u_reports) == 0

    # Temporary analysis does NOT appear in admin reports
    client.post("/api/logout")
    login_admin(client)
    admin_reports = client.get("/api/admin/reports").json.get("reports", [])
    assert len(admin_reports) == initial_reports_count


def test_32_temporary_files_cleanup_policy(test_env):
    # Verify clean_temp_uploads runs without error
    db.clean_temp_uploads()


# ---------------- 4. REGRESSION TESTS ----------------

def test_33_image_preprocessing_pipeline(test_env):
    rid = "TEST-R100"
    out_dir = test_env["tmp_path"] / "preproc_out"
    res = damage_detector.analyze(test_env["test_img_path"], out_dir, rid)
    assert "steps" in res
    assert len(res["steps"]) == 4  # s1, s2, s3, s4
    assert (out_dir / f"{rid}_annotated.jpg").exists()


def test_34_opencv_fallback_works_when_model_absent(test_env):
    rid = "TEST-R101"
    out_dir = test_env["tmp_path"] / "fallback_out"
    res = damage_detector.analyze(test_env["test_img_path"], out_dir, rid)
    assert res.get("fallback_active") is True
    assert "heuristic" in res.get("engine", "").lower()
    assert "heuristic" in res.get("warning", "").lower()


def test_35_analytics_work_with_empty_and_populated_data(test_env):
    empty_df = db.reports_df(user_id=999999)
    s_empty = charts.summary(empty_df)
    assert s_empty["total"] == 0
    c_empty = charts.make_chart("types", empty_df)
    assert len(c_empty) > 0  # Produces empty fallback chart png

    # Populate with report
    client = test_env["client"]
    register_user(client, username="xander", email="xander@test.com")
    uid = db.get_user_by_email("xander@test.com")["user_id"]
    login_admin(client)

    with open(test_env["test_img_path"], "rb") as f:
        client.post("/api/admin/detect-damage", data={
            "user_id": uid, "image": (f, "car.jpg")
        }, content_type="multipart/form-data")

    pop_df = db.reports_df()
    s_pop = charts.summary(pop_df)
    assert s_pop["total"] >= 1
    c_pop = charts.make_chart("types", pop_df)
    assert len(c_pop) > 0


def test_36_excel_export_does_not_corrupt_data(test_env):
    excel_stream = db.export_admin_excel()
    assert excel_stream is not None
    wb = openpyxl.load_workbook(excel_stream)
    assert set(wb.sheetnames) == {"Reports", "Users", "LoginHistory"}


def test_37_invalid_uploads_rejected(test_env):
    client = test_env["client"]
    register_user(client, username="yasmine", email="yasmine@test.com")

    # Invalid extension (.txt)
    text_file = io.BytesIO(b"not an image")
    res = client.post("/api/user/analyze-temporary", data={
        "image": (text_file, "notes.txt")
    }, content_type="multipart/form-data")
    assert res.status_code == 400
    assert "supported" in res.json.get("error", "").lower()


def test_38_unauthorized_api_requests_rejected(test_env):
    client = test_env["client"]
    # No session
    res = client.get("/api/user/dashboard-stats")
    assert res.status_code == 401

    res_admin = client.get("/api/admin/dashboard-stats")
    assert res_admin.status_code == 401


# ---------------- 5. COMPREHENSIVE SECURITY & AUDIT TESTS ----------------

def test_39_admin_password_change_persists_across_restart(test_env):
    """Verifies that updating the admin password in the Admin Profile persists

    across application restarts and does not silently revert.
    """
    client = test_env["client"]
    # 1. Login with initial dev credentials
    res = login_admin(client)
    assert res.status_code == 200

    # 2. Change password via API
    new_pw = "NewAdminSecurePass@2026"
    res_change = client.post("/api/admin/change-password", json={
        "old_password": "Samrat@2026",
        "new_password": new_pw,
        "confirm_password": new_pw
    })
    assert res_change.status_code == 200
    assert res_change.json.get("ok") is True

    # 3. Simulate application restart by re-initializing database
    db.init_db(config.DB_PATH)

    # 4. Old password must be rejected
    res_old = client.post("/api/admin/login", json={
        "email": "admin@Samrat.com",
        "password": "Samrat@2026"
    })
    assert res_old.status_code == 401

    # 5. New password must be accepted
    res_new = client.post("/api/admin/login", json={
        "email": "admin@Samrat.com",
        "password": new_pw
    })
    assert res_new.status_code == 200
    assert res_new.json.get("ok") is True


def test_40_csrf_protection_enforces_tokens_on_state_changing_requests(test_env):
    """Verifies that CSRF protection rejects state-changing requests lacking valid tokens."""
    client = test_env["client"]
    app.config["ENFORCE_CSRF_IN_TESTS"] = True

    try:
        # Request home page to initiate session and generate CSRF token
        client.get("/")
        with client.session_transaction() as sess:
            csrf_token = sess.get("_csrf_token")
        assert csrf_token is not None

        # State-changing request WITHOUT token must fail with 403
        res_no_token = client.post("/api/admin/login", json={
            "email": "admin@Samrat.com",
            "password": "Samrat@2026"
        })
        assert res_no_token.status_code == 403
        assert "CSRF validation failed" in res_no_token.json.get("error", "")

        # State-changing request WITH valid token must succeed
        res_with_token = client.post("/api/admin/login", json={
            "email": "admin@Samrat.com",
            "password": "Samrat@2026"
        }, headers={"X-CSRF-Token": csrf_token})
        assert res_with_token.status_code == 200
        assert res_with_token.json.get("ok") is True

    finally:
        app.config["ENFORCE_CSRF_IN_TESTS"] = False


def test_41_local_path_analysis_rejected_when_disabled(test_env):
    """Verifies that local filesystem path analysis is denied when ALLOW_LOCAL_PATHS is false."""
    client = test_env["client"]
    register_user(client, username="victor_user", email="victor@test.com")
    uid = db.get_user_by_email("victor@test.com")["user_id"]
    login_admin(client)

    # config.ALLOW_LOCAL_PATHS is False by default
    assert config.ALLOW_LOCAL_PATHS is False

    res = client.post("/api/admin/detect-damage", data={
        "user_id": uid,
        "folder": str(test_env["tmp_path"]),
        "name": "test_car.jpg"
    })
    assert res.status_code == 400
    assert "disabled" in res.json.get("error", "").lower()



def test_42_production_security_config_requires_strong_secret_key(monkeypatch):
    """Verifies that production mode strictly requires a secure, non-default SECRET_KEY."""
    import importlib
    monkeypatch.setenv("FLASK_ENV", "production")
    monkeypatch.delenv("SECRET_KEY", raising=False)

    with pytest.raises(RuntimeError, match="CRITICAL SECURITY CONFIGURATION ERROR"):
        # Reloading config under production without SECRET_KEY must raise RuntimeError
        importlib.reload(config)

    # Restoring environment
    monkeypatch.undo()
    importlib.reload(config)


def test_43_unauthenticated_upload_access_rejected(test_env):
    """Verifies that accessing uploaded damage images requires authentication."""
    client = test_env["client"]
    # Without logging in
    res = client.get("/uploads/R-999999_annotated.jpg")
    assert res.status_code == 401
    assert "sign in" in res.json.get("error", "").lower()



def test_44_path_traversal_on_upload_routes_prevented(test_env):
    """Verifies that directory traversal payloads in image paths are safely rejected."""
    client = test_env["client"]
    login_admin(client)

    res = client.get("/uploads/..%2F..%2Fbackend%2Fconfig.py")
    assert res.status_code in (400, 403, 404)

    res_temp = client.get("/uploads/temp/..%2F..%2Fbackend%2Fconfig.py")
    assert res_temp.status_code in (400, 403, 404)


def test_45_admin_download_excel_rejected_for_unauthenticated_and_regular_user(test_env):
    """Verifies that non-admin clients receive 401/403 when trying to download the Excel database."""
    client = test_env["client"]

    # 1. Unauthenticated request -> 401
    res_unauth = client.get("/api/admin/download-excel")
    assert res_unauth.status_code == 401

    # 2. Regular user request -> 403
    register_user(client, username="norman", email="norman@test.com")
    res_user = client.get("/api/admin/download-excel")
    assert res_user.status_code == 403

